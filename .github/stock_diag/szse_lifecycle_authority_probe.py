from __future__ import annotations

import argparse
from datetime import date, datetime
import hashlib
from io import BytesIO
import json
from pathlib import Path
from urllib.parse import urlencode, urlparse
from urllib.request import Request, urlopen

from openpyxl import load_workbook

ENDPOINT = "https://www.szse.cn/api/report/ShowReport"
CURRENT_PARAMS = {"SHOWTYPE": "xlsx", "CATALOGID": "1110", "TABKEY": "tab1", "random": "0.3141592653589793"}
DELIST_PARAMS = {"SHOWTYPE": "xlsx", "CATALOGID": "1793_ssgs", "TABKEY": "tab2", "random": "0.2718281828459045"}
HEADERS = {
    "Referer": "https://www.szse.cn/market/product/stock/list/index.html",
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/154 Safari/537.36",
    "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream,*/*",
}


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def fetch(params: dict[str, str]) -> tuple[bytes, dict]:
    url = ENDPOINT + "?" + urlencode(params)
    req = Request(url, headers=HEADERS, method="GET")
    with urlopen(req, timeout=35) as response:
        raw = response.read()
        final_url = response.geturl()
        status = int(getattr(response, "status", 0) or 0)
        content_type = str(response.headers.get("Content-Type") or "")
    host = (urlparse(final_url).hostname or "").lower()
    if status != 200:
        raise RuntimeError(f"SZSE HTTP {status}: {final_url}")
    if host != "www.szse.cn" and not host.endswith(".szse.cn"):
        raise RuntimeError(f"non-SZSE redirect: {final_url}")
    if not raw.startswith(b"PK"):
        raise RuntimeError(f"SZSE response is not XLSX ZIP bytes; prefix={raw[:16]!r}")
    return raw, {
        "request_url": url,
        "resolved_url": final_url,
        "http_status": status,
        "content_type": content_type,
        "bytes": len(raw),
        "sha256": sha256_bytes(raw),
        "official_host": True,
    }


def norm_code(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        raise ValueError("boolean security code")
    if isinstance(value, int):
        s = str(value)
    elif isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f"non-integer security code {value!r}")
        s = str(int(value))
    else:
        s = str(value).strip()
        if s.endswith(".0") and s[:-2].isdigit():
            s = s[:-2]
    if not s:
        return ""
    if not s.isdigit():
        raise ValueError(f"non-digit security code {s!r}")
    s = s.zfill(6)
    if len(s) != 6:
        raise ValueError(f"invalid six-digit security code {s!r}")
    return s


def norm_date(value: object, *, allow_blank: bool) -> str | None:
    if value is None or str(value).strip() in {"", "-", "None", "nan", "NaT"}:
        if allow_blank:
            return None
        raise ValueError("required date missing")
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        # openpyxl normally converts date-formatted Excel serials, so a bare numeric
        # date is ambiguous and must fail closed instead of guessing.
        raise ValueError(f"ambiguous numeric date {value!r}")
    s = str(value).strip().replace("/", "-").replace(".", "-")
    compact = s.replace("-", "")
    if len(compact) == 8 and compact.isdigit():
        d = date(int(compact[:4]), int(compact[4:6]), int(compact[6:8]))
        return d.isoformat()
    try:
        return date.fromisoformat(s[:10]).isoformat()
    except Exception as exc:
        raise ValueError(f"invalid date {value!r}") from exc


def rows_with_headers(raw: bytes, required: tuple[str, ...]) -> tuple[list[dict], dict]:
    wb = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    if not wb.worksheets:
        raise RuntimeError("XLSX has no worksheets")
    ws = wb.worksheets[0]
    header_row = None
    headers: list[str] = []
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20), values_only=True), start=1):
        vals = [str(v).strip() if v is not None else "" for v in row]
        if all(name in vals for name in required):
            header_row = idx
            headers = vals
            break
    if header_row is None:
        raise RuntimeError(f"required headers not found: {required}; first sheet={ws.title!r}")

    seen = {}
    for i, h in enumerate(headers):
        if h and h not in seen:
            seen[h] = i
    out: list[dict] = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        obj = {name: row[seen[name]] if seen[name] < len(row) else None for name in required}
        if all(v is None or str(v).strip() == "" for v in obj.values()):
            continue
        out.append(obj)
    return out, {
        "sheet": ws.title,
        "header_row": header_row,
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "required_headers": list(required),
    }


def parse_current(raw: bytes) -> tuple[list[dict], dict]:
    required = ("A股代码", "A股简称", "A股上市日期")
    rows, meta = rows_with_headers(raw, required)
    out = []
    for item in rows:
        code = norm_code(item["A股代码"])
        if not code:
            continue
        listed = norm_date(item["A股上市日期"], allow_blank=False)
        out.append({
            "market": "SZSE",
            "code": code,
            "listing_date": listed,
            "delisting_date": None,
            "source_state": "currently_listed",
        })
    if not out:
        raise RuntimeError("current SZSE A-share list parsed zero securities")
    return out, {**meta, "parsed_count": len(out)}


def parse_delisted(raw: bytes) -> tuple[list[dict], dict]:
    required = ("证券代码", "上市日期", "终止上市日期")
    rows, meta = rows_with_headers(raw, required)
    out = []
    for item in rows:
        code = norm_code(item["证券代码"])
        if not code:
            continue
        listed = norm_date(item["上市日期"], allow_blank=False)
        delisted = norm_date(item["终止上市日期"], allow_blank=False)
        if delisted <= listed:
            raise RuntimeError(f"SZSE {code} delisting date is not after listing date")
        out.append({
            "market": "SZSE",
            "code": code,
            "listing_date": listed,
            "delisting_date": delisted,
            "source_state": "terminated_listing",
        })
    if not out:
        raise RuntimeError("SZSE terminated-listing workbook parsed zero securities")
    return out, {**meta, "parsed_count": len(out)}


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--output", required=True, type=Path)
    ap.add_argument("--raw-dir", required=True, type=Path)
    args = ap.parse_args()
    args.raw_dir.mkdir(parents=True, exist_ok=True)

    current_raw, current_transport = fetch(CURRENT_PARAMS)
    delisted_raw, delisted_transport = fetch(DELIST_PARAMS)
    current_path = args.raw_dir / "szse_a_share_current.xlsx"
    delisted_path = args.raw_dir / "szse_terminated_listings.xlsx"
    current_path.write_bytes(current_raw)
    delisted_path.write_bytes(delisted_raw)

    current, current_parse = parse_current(current_raw)
    delisted, delisted_parse = parse_delisted(delisted_raw)

    lifecycle: dict[str, dict] = {}
    conflicts = []
    duplicates = []
    for row in current + delisted:
        code = row["code"]
        previous = lifecycle.get(code)
        if previous is None:
            lifecycle[code] = row
        elif previous == row:
            duplicates.append(code)
        else:
            conflicts.append({"code": code, "first": previous, "second": row})
    if conflicts:
        raise RuntimeError(f"conflicting SZSE lifecycle rows: {conflicts[:5]}")

    ordered = [lifecycle[k] for k in sorted(lifecycle)]
    payload = "".join(
        f"{r['code']}|{r['listing_date']}|{r['delisting_date'] or ''}\n"
        for r in ordered
    ).encode("utf-8")
    projection_path = args.raw_dir / "szse_lifecycle_projection.txt"
    projection_path.write_bytes(payload)

    status = "PASS" if current and delisted and not conflicts else "NOT VERIFIED"
    evidence = {
        "schema": "stock-szse-native-lifecycle-authority-diagnostic-v1",
        "status": status,
        "product_bytes_changed": False,
        "scope": "SZSE native lifecycle source/projection diagnostic only",
        "authority": {
            "provider": "Shenzhen Stock Exchange",
            "endpoint": ENDPOINT,
            "current_a_share": {
                "params": CURRENT_PARAMS,
                "transport": current_transport,
                "raw_path": current_path.as_posix(),
                "parse": current_parse,
            },
            "terminated_listings": {
                "params": DELIST_PARAMS,
                "transport": delisted_transport,
                "raw_path": delisted_path.as_posix(),
                "parse": delisted_parse,
            },
        },
        "lifecycle": {
            "current_count": len(current),
            "terminated_count": len(delisted),
            "union_unique_security_count": len(ordered),
            "duplicate_exact_count": len(duplicates),
            "conflict_count": len(conflicts),
            "projection_path": projection_path.as_posix(),
            "projection_sha256": sha256_bytes(payload),
            "normalization": "current A-share workbook union terminated-listing workbook; exact code/list/delist projection",
        },
        "adjudication": {
            "szse_source_identity": "PASS" if status == "PASS" else "NOT VERIFIED",
            "szse_native_raw_binding": "PASS" if status == "PASS" else "NOT VERIFIED",
            "szse_lifecycle_projection": "SUPPORTED" if status == "PASS" else "NOT VERIFIED",
            "szse_historical_pit_availability": "NOT VERIFIED",
            "bse_native_lifecycle_authority": "NOT VERIFIED",
            "full_three_market_authority": "NOT VERIFIED",
            "full_historical_st_pit": "NOT VERIFIED",
            "advanced_business": "NOT VERIFIED",
            "final_gate": "FAIL",
        },
        "boundary": (
            "PASS proves current first-party SZSE XLSX bytes for the A-share list and terminated-listing "
            "table were fetched, preserved, schema-checked and deterministically projected into a lifecycle "
            "union. It does not prove those fields were historically knowable at each past date, does not "
            "establish BSE authority, and cannot promote full historical ST/PIT or Advanced Business."
        ),
    }
    args.output.write_text(json.dumps(evidence, ensure_ascii=False, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "schema": evidence["schema"],
        "status": evidence["status"],
        "current_count": len(current),
        "terminated_count": len(delisted),
        "union_unique_security_count": len(ordered),
        "conflict_count": len(conflicts),
        "projection_sha256": evidence["lifecycle"]["projection_sha256"],
        "adjudication": evidence["adjudication"],
    }, ensure_ascii=False, sort_keys=True))
    return 0 if status == "PASS" else 2


if __name__ == "__main__":
    raise SystemExit(main())
